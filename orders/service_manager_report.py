"""
Отчёт по замерам для расчёта зарплаты сервис-менеджеров — по образцу
«Отчет по замерам.xlsx».

По каждому выполненному за период замеру:
- количество проёмов — все проёмы плюс все панели в них;
- сумма по замеру — по шкале от количества проёмов (OPENINGS_TARIFF ниже);
- удалённость — км × 30 ₽, если она не оплачена клиентом на месте;
- итого по замеру = сумма по замеру + удалённость.
Итого за месяц — сумма «итого» всех замеров СМ. Номера замеров в отчёте
каждого СМ идут с 1.
"""
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from .api_views import get_measurements_queryset_for_user
from .designer_report import REPORT_ROLES, _date, _ids, _user_name
from .models import DistancePayment

# Сумма по замеру от количества проёмов (проёмы + панели): (до N проёмов включительно, ₽)
OPENINGS_TARIFF = [
    (5, Decimal(750)),
    (10, Decimal(900)),
    (15, Decimal(1200)),
    (20, Decimal(1800)),
    (30, Decimal(2400)),
    (40, Decimal(3600)),
    (60, Decimal(4800)),
]
# Удалённость, ₽ за км (если не оплачена на месте)
DISTANCE_RATE = Decimal(30)


def measurement_sum(openings):
    """Сумма по замеру по шкале. Больше 60 проёмов шкала не описывает — берём верхнюю ступень и помечаем."""
    if openings <= 0:
        return Decimal(0), False
    for limit, amount in OPENINGS_TARIFF:
        if openings <= limit:
            return amount, False
    return OPENINGS_TARIFF[-1][1], True


def payment_status(distance_payment):
    """Столбец «Оплачен/не оплачен» — про удалённость: оплаченная на месте в расчёт не идёт."""
    if distance_payment == DistancePayment.ON_SITE:
        return 'Оплачен на месте'
    if distance_payment == DistancePayment.INVOICE:
        return 'Не оплачен'
    return ''


def _money(value):
    return f'{value:.2f}'


def build_report(request):
    """Строки отчёта, сгруппированные по СМ. Суммы — Decimal (в JSON и Excel приводятся отдельно)."""
    # Видимость — как у раздела «Замеры»: руководителю его город, админу всё.
    # Проёмы, вложения и прочее, что подгружает базовый запрос, отчёту не нужны.
    qs = (
        get_measurements_queryset_for_user(request.user)
        .prefetch_related(None)
        .filter(is_done=True, done_at__isnull=False, service_manager__isnull=False)
    )
    # В расчёт идут замеры, выполненные в выбранном периоде
    date_from, date_to = _date(request, 'date_from'), _date(request, 'date_to')
    if date_from:
        qs = qs.filter(done_at__date__gte=date_from)
    if date_to:
        qs = qs.filter(done_at__date__lte=date_to)
    if cities := _ids(request, 'city'):
        qs = qs.filter(request__order__salon__city_id__in=cities)
    if sms := _ids(request, 'service_manager'):
        qs = qs.filter(service_manager_id__in=sms)

    # Каждая панель в проёме («галочка» + количество) в расчёте — как ещё один проём
    measurements = qs.annotate(
        openings_count=Count('openings', distinct=True),
        panels_count=Sum('openings__panels_count', filter=Q(openings__has_panels=True)),
    ).order_by('done_at', 'id')

    groups = {}
    for m in measurements:
        sm = m.service_manager
        g = groups.get(sm.id)
        if g is None:
            g = groups[sm.id] = {
                'service_manager': {'id': sm.id, 'full_name': _user_name(sm)},
                'measurements_count': 0,
                'openings_total': 0,
                'measurements_sum': Decimal(0),
                'distance_sum': Decimal(0),
                'total': Decimal(0),
                'measurements': [],
            }
        order = m.request.order
        panels = m.panels_count or 0
        openings_total = m.openings_count + panels
        base, tariff_exceeded = measurement_sum(openings_total)
        counted_km = m.distance_km if m.distance_payment == DistancePayment.INVOICE and m.distance_km else Decimal(0)
        distance = counted_km * DISTANCE_RATE
        total = base + distance

        g['measurements_count'] += 1
        g['openings_total'] += openings_total
        g['measurements_sum'] += base
        g['distance_sum'] += distance
        g['total'] += total
        g['measurements'].append({
            # Порядковый номер в отчёте этого СМ — с 1
            'number': g['measurements_count'],
            'id': m.id,
            'order_id': order.id,
            'done_at': m.done_at,
            'address': order.address,
            'client_name': order.client_name,
            'manager_name': _user_name(order.manager),
            'payment_status': payment_status(m.distance_payment),
            'openings_count': m.openings_count,
            'panels_count': panels,
            'openings_total': openings_total,
            'tariff_exceeded': tariff_exceeded,
            'measurement_sum': base,
            'distance_km': m.distance_km,
            'distance_sum': distance,
            'total': total,
        })
    return sorted(groups.values(), key=lambda g: g['service_manager']['full_name'])


def _json_row(row):
    return {
        **row,
        'measurement_sum': _money(row['measurement_sum']),
        'distance_km': f"{row['distance_km']:.1f}" if row['distance_km'] is not None else None,
        'distance_sum': _money(row['distance_sum']),
        'total': _money(row['total']),
    }


class ServiceManagerReportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if request.user.role not in REPORT_ROLES:
            return Response({'detail': 'Отчёт доступен руководителю.'}, status=status.HTTP_403_FORBIDDEN)

        groups = build_report(request)
        if request.query_params.get('export') == 'xlsx':
            return _xlsx_response(groups, request)

        rows = [{
            **g,
            'measurements_sum': _money(g['measurements_sum']),
            'distance_sum': _money(g['distance_sum']),
            'total': _money(g['total']),
            'measurements': [_json_row(r) for r in g['measurements']],
        } for g in groups]
        return Response({
            'totals': {
                'service_managers_count': len(groups),
                'measurements_count': sum(g['measurements_count'] for g in groups),
                'openings_total': sum(g['openings_total'] for g in groups),
                'measurements_sum': _money(sum((g['measurements_sum'] for g in groups), Decimal(0))),
                'distance_sum': _money(sum((g['distance_sum'] for g in groups), Decimal(0))),
                'total': _money(sum((g['total'] for g in groups), Decimal(0))),
            },
            'tariff': {
                'openings': [{'up_to': limit, 'amount': _money(amount)} for limit, amount in OPENINGS_TARIFF],
                'distance_rate': _money(DISTANCE_RATE),
            },
            'service_managers': rows,
        })


XLSX_HEADERS = [
    '№', 'Дата', 'Адрес', 'Номер заказа', 'Менеджер', 'Оплачен/не оплачен',
    'Количество проёмов', 'Сумма по замеру', 'Удалённость', 'Итого по замеру',
]
XLSX_WIDTHS = [6, 12, 40, 14, 22, 20, 12, 16, 14, 16]


def _xlsx_response(groups, request):
    """Отчёт в Excel — по листу на каждого СМ, столбцы как в образце «Отчет по замерам.xlsx»."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    period = ' — '.join(filter(None, [
        request.query_params.get('date_from'), request.query_params.get('date_to'),
    ])) or 'всё время'
    wb = Workbook()
    wb.remove(wb.active)
    used_titles = set()
    for g in groups or [None]:
        name = g['service_manager']['full_name'] if g else 'Нет замеров'
        # Имя листа: до 31 символа и без повторов
        title, n = name[:31], 2
        while title in used_titles:
            title = f'{name[:27]} ({n})'
            n += 1
        used_titles.add(title)
        ws = wb.create_sheet(title)
        ws.append([f'Отчёт по замерам: {name}, период {period}'])
        ws['A1'].font = Font(bold=True)
        ws.append(XLSX_HEADERS)
        for cell in ws[2]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        for row in (g['measurements'] if g else []):
            ws.append([
                row['number'],
                timezone.localtime(row['done_at']).date(),
                row['address'],
                row['order_id'],
                row['manager_name'],
                row['payment_status'],
                row['openings_total'],
                float(row['measurement_sum']),
                float(row['distance_sum']),
                float(row['total']),
            ])
            ws.cell(ws.max_row, 2).number_format = 'DD.MM.YYYY'
        if g:
            ws.append(['Итого за месяц', None, None, None, None, None, g['openings_total'],
                       float(g['measurements_sum']), float(g['distance_sum']), float(g['total'])])
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
        for i, width in enumerate(XLSX_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for r in ws.iter_rows(min_row=3, min_col=8, max_col=10):
            for cell in r:
                cell.number_format = '#,##0.00'

    buf = BytesIO()
    wb.save(buf)
    stamp = datetime.now().strftime('%Y%m%d')
    response = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="sm_report_{stamp}.xlsx"'
    return response
