"""
Отчёт по замерам для расчёта зарплаты СМ (по образцу «Отчет по замерам.xlsx»):
шкала по количеству проёмов, удалённость × 30 ₽, итоги, выгрузка в Excel.

Запуск (роль Postgres не умеет создавать БД, поэтому через sqlite):
    DATABASE_ENGINE=django.db.backends.sqlite3 DATABASE_NAME=/tmp/t.sqlite3 \\
        python manage.py test orders.tests_service_manager_report
"""
from datetime import datetime
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from orders.models import Measurement, MeasurementOpening, MeasurementRequest, Order, OrderStatus, Salon
from users.models import City

User = get_user_model()
URL = '/api/v1/reports/service-managers/'


def aware(y, m, d):
    return timezone.make_aware(datetime(y, m, d, 12, 0))


class ServiceManagerReportTest(TestCase):
    def setUp(self):
        self.kazan = City.objects.create(name='Казань')
        self.samara = City.objects.create(name='Самара')
        self.salon = Salon.objects.create(name='Академи', city=self.kazan)
        self.far_salon = Salon.objects.create(name='Самарский', city=self.samara)
        self.leader = User.objects.create_user(username='boss', password='x', role='leader', city=self.kazan)
        self.manager = User.objects.create_user(
            username='mgr', password='x', role='manager', city=self.kazan, salon=self.salon,
        )
        self.far_manager = User.objects.create_user(
            username='far', password='x', role='manager', city=self.samara, salon=self.far_salon,
        )
        self.petr = User.objects.create_user(
            username='petr', password='x', role='service_manager', city=self.kazan,
            first_name='Пётр', last_name='Петров',
        )
        self.ivan = User.objects.create_user(
            username='ivan', password='x', role='service_manager', city=self.kazan,
            first_name='Иван', last_name='Иванов',
        )
        self.far_sm = User.objects.create_user(username='farsm', password='x', role='service_manager', city=self.samara)

        self.m1 = self.measurement(self.petr, aware(2026, 9, 10), openings=3)
        self.m2 = self.measurement(self.petr, aware(2026, 10, 2))
        self.m3 = self.measurement(self.ivan, aware(2026, 10, 5))
        self.measurement(self.petr, None)                                 # не выполнен
        self.measurement(self.far_sm, aware(2026, 10, 3), manager=self.far_manager)  # чужой город

        self.client = APIClient()
        self.client.force_authenticate(self.leader)

    def measurement(self, sm, done_at, openings=0, manager=None):
        manager = manager or self.manager
        order = Order.objects.create(
            manager=manager, salon=manager.salon, client_name=f'Клиент {sm.username}',
            address='ул. Ленина, 1', kp_number='КП-1', status=OrderStatus.MEASUREMENT_DONE,
        )
        req = MeasurementRequest.objects.create(order=order, contact_name='К', contact_phone='+7', created_by=manager)
        m = Measurement.objects.create(request=req, service_manager=sm, is_done=bool(done_at), done_at=done_at)
        for n in range(openings):
            MeasurementOpening.objects.create(measurement=m, opening_number=n + 1)
        return m

    def get(self, **params):
        response = self.client.get(URL, params)
        self.assertEqual(response.status_code, 200, response.data)
        return response.data

    def counts(self, data):
        return {r['service_manager']['id']: r['measurements_count'] for r in data['service_managers']}

    def row(self, data, sm):
        return next(r for r in data['service_managers'] if r['service_manager']['id'] == sm.id)

    def test_counts_done_measurements_in_own_city(self):
        data = self.get()
        self.assertEqual(self.counts(data), {self.petr.id: 2, self.ivan.id: 1})
        # m1 — 3 проёма → 750 ₽; у m2 и m3 проёмов нет → 0
        self.assertEqual(data['totals']['measurements_count'], 3)
        self.assertEqual(data['totals']['total'], '750.00')

    def test_period_is_by_done_date(self):
        data = self.get(date_from='2026-10-01', date_to='2026-10-31')
        self.assertEqual(self.counts(data), {self.petr.id: 1, self.ivan.id: 1})

    def test_measurement_row_like_template(self):
        petr = self.row(self.get(), self.petr)
        self.assertEqual(petr['service_manager']['full_name'], 'Пётр Петров')
        first, second = petr['measurements']
        # Номера в отчёте СМ — с 1, по дате выполнения
        self.assertEqual((first['number'], second['number']), (1, 2))
        self.assertEqual(first['id'], self.m1.id)
        self.assertEqual(first['address'], 'ул. Ленина, 1')
        self.assertEqual(first['order_id'], self.m1.request.order_id)
        self.assertEqual(first['manager_name'], 'mgr')
        self.assertEqual(first['openings_total'], 3)
        self.assertEqual(first['measurement_sum'], '750.00')
        self.assertEqual(first['total'], '750.00')

    def test_filter_by_service_manager(self):
        self.assertEqual(self.counts(self.get(service_manager=self.ivan.id)), {self.ivan.id: 1})

    def test_service_manager_sees_only_own_report(self):
        self.client.force_authenticate(self.petr)
        # Даже если попросить чужого СМ — отдаём только свой отчёт
        data = self.get(service_manager=self.ivan.id)
        self.assertEqual(self.counts(data), {self.petr.id: 2})

    def test_manager_has_no_access(self):
        self.client.force_authenticate(self.manager)
        self.assertEqual(self.client.get(URL).status_code, 403)

    def test_distance_counted_only_if_not_paid_on_site(self):
        Measurement.objects.filter(pk=self.m1.pk).update(distance_payment='invoice', distance_km=Decimal('12.5'))
        Measurement.objects.filter(pk=self.m3.pk).update(distance_payment='on_site', distance_km=Decimal('50'))
        data = self.get()
        first = self.row(data, self.petr)['measurements'][0]
        self.assertEqual(first['payment_status'], 'Не оплачен')
        self.assertEqual(first['distance_sum'], '375.00')      # 12,5 км × 30
        self.assertEqual(first['total'], '1125.00')            # 750 + 375
        ivan = self.row(data, self.ivan)
        self.assertEqual(ivan['measurements'][0]['payment_status'], 'Оплачен на месте')
        self.assertEqual(ivan['measurements'][0]['distance_sum'], '0.00')
        self.assertEqual(data['totals']['total'], '1125.00')

    def test_panels_count_as_extra_openings(self):
        # В m1 три проёма: у двух галочка «Панели» (2 и 1 шт.), у третьего количество без галочки не в счёт
        o1, o2, o3 = self.m1.openings.order_by('opening_number')
        MeasurementOpening.objects.filter(pk=o1.pk).update(has_panels=True, panels_count=2)
        MeasurementOpening.objects.filter(pk=o2.pk).update(has_panels=True, panels_count=1)
        MeasurementOpening.objects.filter(pk=o3.pk).update(has_panels=False, panels_count=5)
        first = self.row(self.get(), self.petr)['measurements'][0]
        self.assertEqual((first['openings_count'], first['panels_count'], first['openings_total']), (3, 3, 6))
        self.assertEqual(first['measurement_sum'], '900.00')   # 6 проёмов → ступень 6–10

    def test_month_total_is_sum_of_measurement_totals(self):
        for n in range(4, 13):                                 # m1 → 12 проёмов: 1200 ₽
            MeasurementOpening.objects.create(measurement=self.m1, opening_number=n)
        Measurement.objects.filter(pk=self.m2.pk).update(distance_payment='invoice', distance_km=Decimal('10'))
        petr = self.row(self.get(), self.petr)
        self.assertEqual(petr['measurements_sum'], '1200.00')
        self.assertEqual(petr['distance_sum'], '300.00')
        self.assertEqual(petr['total'], '1500.00')

    def test_excel_export(self):
        from io import BytesIO
        from openpyxl import load_workbook
        response = self.client.get(URL, {'export': 'xlsx'})
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        ws = wb['Пётр Петров']
        self.assertEqual(ws['A2'].value, '№')
        self.assertEqual(ws['J2'].value, 'Итого по замеру')
        self.assertEqual(ws['A3'].value, 1)
        self.assertEqual(ws['J3'].value, 750)
        self.assertEqual(ws['A5'].value, 'Итого за месяц')
        self.assertEqual(ws['J5'].value, 750)


class TariffTest(SimpleTestCase):
    """Шкала «сумма по замеру» из образца отчёта."""

    def test_tariff_steps(self):
        from orders.service_manager_report import measurement_sum
        cases = {0: 0, 1: 750, 5: 750, 6: 900, 10: 900, 11: 1200, 15: 1200, 16: 1800, 20: 1800,
                 21: 2400, 30: 2400, 31: 3600, 40: 3600, 41: 4800, 60: 4800}
        for openings, amount in cases.items():
            self.assertEqual(measurement_sum(openings), (Decimal(amount), False), openings)

    def test_over_sixty_is_flagged(self):
        from orders.service_manager_report import measurement_sum
        self.assertEqual(measurement_sum(61), (Decimal(4800), True))
