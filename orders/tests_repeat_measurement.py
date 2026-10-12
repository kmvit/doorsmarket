"""
Повторный замер — новый замер рядом с прежним: копия данных, свой номер,
прежний сохраняется; в отчёте по ЗП СМ — оба; связка с позициями по
умолчанию по последнему, менеджер может выбрать другой.

Запуск (роль Postgres не умеет создавать БД, поэтому через sqlite):
    DATABASE_ENGINE=django.db.backends.sqlite3 DATABASE_NAME=/tmp/t.sqlite3 \\
        python manage.py test orders.tests_repeat_measurement
"""
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from orders.models import (
    Measurement, MeasurementAttachment, MeasurementOpening, MeasurementRequest,
    MeasurementSignature, Order, OrderItem, OrderStatus, Salon,
)
from users.models import City

User = get_user_model()


@patch('users.push_utils.send_push_notification')
class RepeatMeasurementTest(TestCase):
    def setUp(self):
        city = City.objects.create(name='Казань')
        salon = Salon.objects.create(name='Салон', city=city)
        self.manager = User.objects.create_user(
            username='mgr', password='x', role='manager', city=city, salon=salon,
        )
        self.sm = User.objects.create_user(username='sm', password='x', role='service_manager', city=city)
        self.order = Order.objects.create(
            manager=self.manager, salon=salon, client_name='Клиент', status=OrderStatus.MEASUREMENT_PROCESSED,
        )
        self.item = OrderItem.objects.create(order=self.order, opening_number=1, model_name='Дверь', quantity=1)
        mr = MeasurementRequest.objects.create(
            order=self.order, contact_name='Клиент', contact_phone='+700', created_by=self.manager,
        )
        self.first = Measurement.objects.create(
            request=mr, service_manager=self.sm, is_done=True, done_at=timezone.now(),
            is_processed=True, processed_at=timezone.now(), measurement_date=timezone.now(),
            distance_km=12, paid_on_site=True,
        )
        self.opening = MeasurementOpening.objects.create(
            measurement=self.first, opening_number=1, room_name='Спальня',
            actual_height=2050, actual_width=850, order_item=self.item,
        )
        MeasurementAttachment.objects.create(
            measurement=self.first, opening=self.opening, file='orders/measurements/photo.jpg', name='фото',
        )
        MeasurementSignature.objects.create(measurement=self.first, file='orders/signatures/s.jpg')
        self.client = APIClient()
        self.client.force_authenticate(self.manager)

    def _repeat(self, m=None):
        return self.client.post(
            f'/api/v1/measurements/{(m or self.first).id}/request_repeat/',
            {'reason': 'Клиент перенёс стену'}, format='json',
        )

    def test_repeat_creates_filled_copy_and_keeps_first(self, _push):
        response = self._repeat()
        self.assertEqual(response.status_code, 200, response.data)
        new = Measurement.objects.get(pk=response.data['id'])

        self.assertNotEqual(new.id, self.first.id)
        self.assertEqual(new.previous_id, self.first.id)
        self.assertEqual(new.service_manager_id, self.sm.id)
        self.assertFalse(new.is_done)
        self.assertIsNone(new.measurement_date)
        self.assertEqual(new.distance_km, 12)
        self.assertTrue(new.paid_on_site)
        op = new.openings.get()
        self.assertEqual((op.room_name, op.actual_height, op.order_item_id), ('Спальня', 2050, self.item.id))
        self.assertEqual(new.attachments.get().opening_id, op.id)
        self.assertEqual(new.signatures.count(), 0)

        self.first.refresh_from_db()
        self.assertTrue(self.first.is_done and self.first.is_processed)
        self.assertEqual(self.first.openings.get().actual_height, 2050)

        self.order.refresh_from_db()
        self.assertEqual(self.order.status, OrderStatus.MEASUREMENT_REQUESTED)

    def test_order_lists_measurements_and_links_to_latest(self, _push):
        new_id = self._repeat().data['id']
        data = self.client.get(f'/api/v1/orders/{self.order.id}/').data
        self.assertEqual([m['id'] for m in data['measurements']], [self.first.id, new_id])
        self.assertEqual(data['link_measurement_id'], new_id)

        response = self.client.post(
            f'/api/v1/orders/{self.order.id}/set_link_measurement/',
            {'measurement_id': self.first.id}, format='json',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['link_measurement_id'], self.first.id)
        item_data = response.data['items'][0]['measurement_data']
        self.assertEqual(item_data['opening_id'], self.opening.id)

    def test_repeat_only_from_latest_and_old_not_processable(self, _push):
        self._repeat()
        self.assertEqual(self._repeat().status_code, 400)
        response = self.client.post(f'/api/v1/measurements/{self.first.id}/mark_processed/')
        self.assertEqual(response.status_code, 400)

    def test_both_in_service_manager_report(self, _push):
        new = Measurement.objects.get(pk=self._repeat().data['id'])
        new.is_done, new.done_at = True, timezone.now()
        new.save()
        sm_client = APIClient()
        sm_client.force_authenticate(self.sm)
        response = sm_client.get('/api/v1/reports/service-managers/')
        self.assertEqual(response.status_code, 200, response.data)
        rows = response.data['service_managers'][0]['measurements']
        self.assertEqual({r['id'] for r in rows}, {self.first.id, new.id})
        repeat_row = next(r for r in rows if r['id'] == new.id)
        self.assertEqual(repeat_row['previous_id'], self.first.id)

        # В Excel — номер замера, у повторного с пометкой
        from io import BytesIO
        from openpyxl import load_workbook
        xlsx = sm_client.get('/api/v1/reports/service-managers/', {'export': 'xlsx'})
        ws = load_workbook(BytesIO(xlsx.content)).active
        labels = {ws.cell(r, 5).value for r in range(3, 5)}
        self.assertEqual(labels, {self.first.id, f'{new.id} (повторный к №{self.first.id})'})
